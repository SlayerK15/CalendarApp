from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font

from app.excel import read_excel
from app.parser import parse_sheet


def workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "Term-I"
    sheet["A2"] = "Term-I Class Schedule_MBA & MBAA Batch 2026-28"
    sheet["B5"], sheet["C5"] = "Date", "Day"
    for col, section in [(4, "A"), (9, "B"), (14, "C"), (19, "D"), (24, "E"), (29, "F")]:
        analytics = "(MBA-Analytics Section)" if section in "EF" else ""
        sheet.cell(4, col, f"Section - {section} {analytics} (Classroom-A-1 Ground Floor)")
        for index, time in enumerate(
            ["9.15 am -\n10.30 am", "10.45 am-12.00 pm", "12.15 pm-01.30 pm", "02.45 pm-04.00 pm"]
        ):
            sheet.cell(5, col + index, time)
    sheet["B7"], sheet["B8"] = date(2026, 9, 10), date(2026, 9, 11)
    return book, sheet


def parse(book):
    buffer = BytesIO()
    book.save(buffer)
    return parse_sheet(read_excel(buffer.getvalue()))


def test_grid_sections_and_cell_specific_times_rooms_and_reservations():
    book, sheet = workbook()
    sheet["D7"] = "ME-1\nProf. Example"
    sheet["G7"] = "BS-4\nProf. Example\n----------\nOB-Quiz-1\n04:15 pm to 04:35 pm"
    sheet["X7"] = "MF-1\nProf. Example\nCR-C2"
    sheet["AD7"] = "Booked for MBAA Second Year"
    sheet["AF7"] = "Booked for a meeting\n----------\nFA-Quiz-1\n05:00 pm to 05:15 pm"
    sheet["D8"] = "HOLIDAY: NO MBA/MBAA CLASSES"
    events = parse(book)
    assert len(events) == 5
    by_title = {event["payload"]["summary"]: event for event in events}
    assert by_title["ME-1"]["section"] == "A"
    assert by_title["ME-1"]["payload"]["start"]["dateTime"] == "2026-09-10T09:15:00+05:30"
    assert by_title["OB-Quiz-1"]["payload"]["start"]["dateTime"] == "2026-09-10T16:15:00+05:30"
    assert by_title["MF-1"]["programme"] == "MBA Analytics"
    assert by_title["MF-1"]["payload"]["location"] == "Classroom C2"
    assert by_title["FA-Quiz-1"]["section"] == "F"


def test_cell_and_partial_rich_text_strikethrough_cancel_only_affected_class():
    book, sheet = workbook()
    sheet["D7"] = "ME-1\nProf. Example"
    sheet["D7"].font = Font(strike=True)
    sheet["G7"] = CellRichText(
        TextBlock(InlineFont(strike=True), "BS-4\nProf. Example"),
        "\n----------\nOB-Quiz-1\n04:15 pm to 04:35 pm",
    )
    sheet["X7"] = "MF-1\nProf. Example\nCancelled"
    result = parse(book)
    assert {e["payload"]["summary"]: e["cancelled"] for e in result} == {
        "ME-1": True,
        "BS-4": True,
        "OB-Quiz-1": False,
        "MF-1": True,
    }


@pytest.mark.parametrize(
    "time,expected",
    [
        ("04:15 to 04:45 pm", ("16:15", "16:45")),
        ("04:15 pm to\n04:35 pm", ("16:15", "16:35")),
        ("04:15 pm to 04:30", ("16:15", "16:30")),
        ("04:15 pm to 0530 pm", ("16:15", "17:30")),
        ("05:00 pmm to 05:15 pm", ("17:00", "17:15")),
    ],
)
def test_observed_time_notation_variants(time, expected):
    book, sheet = workbook()
    sheet["V7"] = "MM-I\nQuiz-2\n" + time
    payload = parse(book)[0]["payload"]
    assert payload["summary"] == "MM-I-Quiz-2"
    assert payload["start"]["dateTime"][11:16] == expected[0]
    assert payload["end"]["dateTime"][11:16] == expected[1]


def test_merged_class_uses_last_slot_end_and_online_location():
    book, sheet = workbook()
    sheet["D7"] = "OB-1\nProf. Example\n(Online)"
    sheet.merge_cells("D7:E7")
    payload = parse(book)[0]["payload"]
    assert payload["end"]["dateTime"][11:16] == "12:00"
    assert payload["location"] == "Online"


def test_session_identity_survives_rescheduling_and_room_changes():
    book, sheet = workbook()
    sheet["D7"] = "ME-1\nProf. Example"
    initial = parse(book)[0]
    sheet["D7"] = None
    sheet["E8"] = "ME-1\nProf. Other\nCR-B2"
    updated = parse(book)[0]
    assert updated["row_id"] == initial["row_id"]
    assert updated["payload"] != initial["payload"]


def test_repeated_labels_are_unique_and_keep_cancelled_occurrence():
    book, sheet = workbook()
    sheet["D7"] = sheet["E7"] = "OB-5\nProf. Example"
    sheet["D7"].font = Font(strike=True)
    result = parse(book)
    assert len({e["row_id"] for e in result}) == 2
    assert [e["cancelled"] for e in result] == [True, False]


@pytest.mark.parametrize("bad", ["Unknown schedule text", "ME-1\n----------\nME-2"])
def test_ambiguous_cells_stop_sync(bad):
    book, sheet = workbook()
    sheet["D7"] = bad
    with pytest.raises(ValueError):
        parse(book)


def test_invalid_class_date_stops_sync():
    book, sheet = workbook()
    sheet["B7"], sheet["D7"] = "not a date", "ME-1"
    with pytest.raises(ValueError, match="unreadable date"):
        parse(book)


def test_cancelled_replacement_can_share_default_slot():
    book, sheet = workbook()
    sheet["AF7"] = CellRichText(
        TextBlock(InlineFont(strike=True), "FM-8\nProf. Former"),
        "\n---------\nWOC-7\nProf. Replacement\n---------\nWOC-6\nProf. Example\n04:15 pm to 05:30 pm",
    )
    result = parse(book)
    assert [(e["payload"]["summary"], e["cancelled"]) for e in result] == [
        ("FM-8", True),
        ("WOC-7", False),
        ("WOC-6", False),
    ]
    assert [e["payload"]["start"]["dateTime"][11:16] for e in result] == ["14:45", "14:45", "16:15"]


def test_all_cancelled_alternatives_are_preserved():
    book, sheet = workbook()
    sheet["K7"] = "MM-I-10\nProf. Example\n---------\nFA-10\nProf. Other"
    sheet["K7"].font = Font(strike=True)
    assert all(event["cancelled"] for event in parse(book))


def test_extra_double_session_and_roman_numeral_course():
    book, sheet = workbook()
    sheet["G7"] = "WOC-Extra Session-1&2\n02:45 pm to 05:30 pm\n---------\nMM-II-Quiz-2\n05:40 pm to 05:50 pm"
    result = parse(book)
    assert [e["payload"]["summary"] for e in result] == ["WOC-Extra Session-1&2", "MM-II-Quiz-2"]
    assert result[0]["payload"]["end"]["dateTime"][11:16] == "17:30"
