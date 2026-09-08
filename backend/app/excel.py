from datetime import date, datetime, time
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils.cell import range_boundaries

MAX_DOWNLOAD_BYTES = 25 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_CELLS = 250_000


class TimetableRows(list):
    """Cell text plus the layout and cancellation metadata needed by grid readers."""

    def __init__(self):
        super().__init__()
        self.strike_spans = {}
        self.merges = {}


def read_excel(content: bytes, sheet_name: str = ""):
    """Read a single XLSX tab in memory; never save or modify the source file."""
    if len(content) > MAX_DOWNLOAD_BYTES:
        raise ValueError("The Excel timetable is too large to read safely (25 MB limit).")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 5000 or sum(item.file_size for item in entries) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("The Excel timetable expands beyond the supported workbook size.")
            # Bound cell objects before loading rich text/merges into memory.
            allocated = 0
            for item in entries:
                if not item.filename.startswith("xl/worksheets/") or not item.filename.endswith(".xml"):
                    continue
                with archive.open(item) as xml:
                    for _, element in iterparse(xml, events=("end",)):
                        tag = element.tag.rsplit("}", 1)[-1]
                        if tag == "c":
                            allocated += 1
                        elif tag == "mergeCell":
                            left, top, right, bottom = range_boundaries(element.attrib["ref"])
                            allocated += (right - left + 1) * (bottom - top + 1)
                        if allocated > MAX_CELLS:
                            raise ValueError("The Excel workbook exceeds the supported cell count.")
                        element.clear()
        workbook = load_workbook(BytesIO(content), data_only=True, keep_links=False, rich_text=True)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("The timetable could not be read as an Excel .xlsx workbook.") from exc
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError("The configured Excel worksheet was not found: " + sheet_name)
            sheet = workbook[sheet_name]
        elif len(workbook.worksheets) == 1:
            sheet = workbook.worksheets[0]
        else:
            raise ValueError(
                "This Excel workbook has multiple worksheets. Set EXCEL_SHEET_NAME to the timetable tab: "
                + ", ".join(workbook.sheetnames[:20])
            )
        if sheet.max_row and sheet.max_column and sheet.max_row * sheet.max_column > MAX_CELLS:
            raise ValueError("The Excel worksheet exceeds the supported size.")
        rows = TimetableRows()
        for merge in sheet.merged_cells.ranges:
            rows.merges[(merge.min_row - 1, merge.min_col - 1)] = (merge.max_row - 1, merge.max_col - 1)
        cells_read = 0
        for cells in sheet.iter_rows():
            cells_read += len(cells)
            if cells_read > MAX_CELLS:
                raise ValueError("The Excel worksheet exceeds the supported size.")
            row = []
            for cell in cells:
                value = cell.value
                if value is not None:
                    raw = str(value)
                    offset = len(raw) - len(raw.lstrip())
                    spans = []
                    if isinstance(value, CellRichText):
                        position = 0
                        for part in value:
                            if isinstance(part, TextBlock) and part.font.strike:
                                spans.append((position - offset, position + len(str(part)) - offset))
                            position += len(str(part))
                    if cell.font.strike:
                        spans = [(0, len(raw.strip()))]
                    if spans:
                        rows.strike_spans[(cell.row - 1, cell.column - 1)] = spans
                if isinstance(value, (datetime, date)):
                    value = value.strftime("%Y-%m-%d")
                elif isinstance(value, time):
                    value = value.strftime("%H:%M")
                row.append("" if value is None else str(value).strip())
            while row and not row[-1]:
                row.pop()
            rows.append(row)
        while rows and not rows[-1]:
            rows.pop()
        return rows
    finally:
        workbook.close()
